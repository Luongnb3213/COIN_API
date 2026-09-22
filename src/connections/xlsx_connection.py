from __future__ import annotations

import threading
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src import config
from src.utils.logger import get_logger


log = get_logger("xlsx_connection")

ACCOUNT_HEADERS = [
    "phone",
    "password",
    "katakana_last_name",
    "katakana_first_name",
    "date_of_birth",
    "pin",
    "status",
    "step_status",
    "error_details",
    "proxy_used",
    "proxy_id",
    "coin_id",
    "customer_status",
    "pcard_status",
    "registered_at",
]
PROXIES_HEADERS = ["proxy", "status", "proxy_id"]
SUCCESS_HEADERS = ACCOUNT_HEADERS

STATUS_VALUES = {"PENDING", "PROCESSING", "SUCCESS", "FAILED", "RETRY", "FAIL_NO_RETRY"}
RUNNABLE_STATUS = {"", "PENDING", "FAILED", "RETRY"}

STATUS_ALIASES = {
    "": "PENDING",
    "pending": "PENDING",
    "processing": "PROCESSING",
    "success": "SUCCESS",
    "ok": "SUCCESS",
    "failed": "FAILED",
    "fail": "FAILED",
    "error": "FAILED",
    "retry": "RETRY",
    "fail_no_retry": "FAIL_NO_RETRY",
    "fatal": "FAIL_NO_RETRY",
}


def normalize_status(value: str) -> str:
    raw = str(value or "").strip()
    if raw.upper() in STATUS_VALUES:
        return raw.upper()
    return STATUS_ALIASES.get(raw.lower(), "FAILED")


class XlsxConnection:
    def __init__(self, xlsx_path: str) -> None:
        self.xlsx_path = Path(xlsx_path).expanduser() if xlsx_path else None
        self._lock = threading.Lock()
        self._connected = False
        if self.xlsx_path and self.xlsx_path.exists():
            self._ensure_workbook()
            self._connected = True
            log.info("Kết nối XLSX: %s", self.xlsx_path)
        else:
            log.warning("File XLSX chưa tồn tại: %s", self.xlsx_path)

    def is_connected(self) -> bool:
        return self._connected

    def _load_workbook(self, read_only: bool = False):
        if not self.xlsx_path:
            raise RuntimeError("Chưa cấu hình xlsx_path.")
        return openpyxl.load_workbook(str(self.xlsx_path), read_only=read_only)

    def _atomic_save(self, wb) -> None:
        if not self.xlsx_path:
            raise RuntimeError("Chưa cấu hình xlsx_path.")
        tmp = self.xlsx_path.with_suffix(self.xlsx_path.suffix + ".tmp")
        wb.save(tmp)
        tmp.replace(self.xlsx_path)

    @staticmethod
    def _headers(ws) -> list[str]:
        return [str(cell.value or "").strip() for cell in ws[1]]

    @staticmethod
    def _col_map(headers: list[str]) -> dict[str, int]:
        return {name: idx + 1 for idx, name in enumerate(headers)}

    def _ensure_workbook(self) -> None:
        with self._lock:
            wb = self._load_workbook()
            if config.ACTIVE_SHEET not in wb.sheetnames:
                ws = wb.create_sheet(config.ACTIVE_SHEET)
                ws.append(ACCOUNT_HEADERS)
            ws = wb[config.ACTIVE_SHEET]
            headers = self._headers(ws)
            if not headers or headers == [""]:
                ws.append(ACCOUNT_HEADERS)
                headers = self._headers(ws)
            for header in ACCOUNT_HEADERS:
                if header not in headers:
                    ws.cell(row=1, column=len(headers) + 1, value=header)
                    headers.append(header)
            self._format_header(ws)
            self._ensure_sheet(wb, "Proxies", PROXIES_HEADERS)
            self._ensure_sheet(wb, "Success", SUCCESS_HEADERS)
            self._atomic_save(wb)
            wb.close()

    def _ensure_sheet(self, wb, name: str, headers: list[str]) -> None:
        if name not in wb.sheetnames:
            ws = wb.create_sheet(name)
            ws.append(headers)
            self._format_header(ws)
            return
        ws = wb[name]
        existing = self._headers(ws)
        if not existing or existing == [""]:
            ws.append(headers)
            existing = self._headers(ws)
        for header in headers:
            if header not in existing:
                ws.cell(row=1, column=len(existing) + 1, value=header)
                existing.append(header)
        self._format_header(ws)

    @staticmethod
    def _format_header(ws) -> None:
        fill = PatternFill("solid", fgColor="1F4E79")
        for cell in ws[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
            cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    def reset_interrupted_to_pending(self) -> None:
        with self._lock:
            wb = self._load_workbook()
            ws = wb[config.ACTIVE_SHEET]
            headers = self._headers(ws)
            col = self._col_map(headers).get("status")
            if not col:
                wb.close()
                return
            changed = 0
            for row in range(2, ws.max_row + 1):
                if normalize_status(ws.cell(row, col).value) == "PROCESSING":
                    ws.cell(row, col, "PENDING")
                    changed += 1
            if changed:
                self._atomic_save(wb)
                log.info("Reset %s dòng PROCESSING về PENDING.", changed)
            wb.close()

    def get_pending_accounts(self, batch_size: int = 50) -> list[dict]:
        with self._lock:
            wb = self._load_workbook(read_only=True)
            ws = wb[config.ACTIVE_SHEET]
            headers = self._headers(ws)
            results: list[dict] = []
            for row_idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                row = {header: values[idx] if idx < len(values) else "" for idx, header in enumerate(headers)}
                if not any(str(row.get(col) or "").strip() for col in ("phone", "password", "katakana_last_name")):
                    continue
                status = normalize_status(str(row.get("status") or ""))
                if status not in RUNNABLE_STATUS:
                    continue
                account = {header: str(row.get(header) or "").strip() for header in ACCOUNT_HEADERS}
                account["_row"] = row_idx
                results.append(account)
                if len(results) >= batch_size:
                    break
            wb.close()
            log.info("Đọc được %s account runnable từ XLSX.", len(results))
            return results

    def update_account(self, row_number: int, data: dict) -> None:
        if row_number < 2:
            return
        with self._lock:
            wb = self._load_workbook()
            ws = wb[config.ACTIVE_SHEET]
            headers = self._headers(ws)
            for header in ACCOUNT_HEADERS:
                if header not in headers:
                    ws.cell(row=1, column=len(headers) + 1, value=header)
                    headers.append(header)
            col_map = self._col_map(headers)
            for key, value in data.items():
                if key not in col_map:
                    continue
                if key == "status":
                    value = normalize_status(value)
                ws.cell(row=row_number, column=col_map[key], value=str(value or ""))
            self._atomic_save(wb)
            wb.close()

    def append_success_account(self, data: dict) -> None:
        phone = str(data.get("phone") or "").strip()
        if not phone:
            return
        with self._lock:
            wb = self._load_workbook()
            self._ensure_sheet(wb, "Success", SUCCESS_HEADERS)
            ws = wb["Success"]
            headers = self._headers(ws)
            col_map = self._col_map(headers)
            phone_col = col_map.get("phone", 1)
            target_row = None
            for row_number in range(2, ws.max_row + 1):
                if str(ws.cell(row=row_number, column=phone_col).value or "").strip() == phone:
                    target_row = row_number
                    break
            if target_row is None:
                target_row = ws.max_row + 1
            for key, value in data.items():
                if key in col_map:
                    ws.cell(row=target_row, column=col_map[key], value=str(value or ""))
            self._atomic_save(wb)
            wb.close()

    def get_active_proxies(self) -> list[dict]:
        with self._lock:
            wb = self._load_workbook()
            if "Proxies" not in wb.sheetnames:
                wb.close()
                return []
            ws = wb["Proxies"]
            headers = self._headers(ws)
            for header in PROXIES_HEADERS:
                if header not in headers:
                    ws.cell(row=1, column=len(headers) + 1, value=header)
                    headers.append(header)
            col_map = self._col_map(headers)
            changed = False
            results = []
            used_ids = set()
            next_id = 1
            for row_number in range(2, ws.max_row + 1):
                raw = str(ws.cell(row=row_number, column=col_map["proxy"]).value or "").strip()
                if not raw:
                    continue
                proxy_id = str(ws.cell(row=row_number, column=col_map["proxy_id"]).value or "").strip()
                if not proxy_id or proxy_id in used_ids:
                    while f"PX-{next_id:06d}" in used_ids:
                        next_id += 1
                    proxy_id = f"PX-{next_id:06d}"
                    next_id += 1
                    ws.cell(row=row_number, column=col_map["proxy_id"], value=proxy_id)
                    changed = True
                used_ids.add(proxy_id)
                status = str(ws.cell(row=row_number, column=col_map["status"]).value or "").strip().lower()
                if status in ("disabled", "inactive", "used", "dead", "0"):
                    continue
                results.append({"proxy_id": proxy_id, "raw": raw, "sheet_row": row_number})
            if changed:
                self._atomic_save(wb)
            wb.close()
            log.info("Đọc được %s proxy ACTIVE từ XLSX.", len(results))
            return results

    def update_proxy_status(self, proxy_ref, status: str = "INACTIVE") -> bool:
        proxy_id = str((proxy_ref or {}).get("proxy_id") or "").strip() if isinstance(proxy_ref, dict) else ""
        raw = str((proxy_ref or {}).get("raw") or proxy_ref or "").strip()
        sheet_row = (proxy_ref or {}).get("sheet_row") if isinstance(proxy_ref, dict) else None
        if not proxy_id and not raw and not sheet_row:
            return False
        with self._lock:
            wb = self._load_workbook()
            if "Proxies" not in wb.sheetnames:
                wb.close()
                return False
            ws = wb["Proxies"]
            headers = self._headers(ws)
            col_map = self._col_map(headers)
            matched = False
            for row_number in range(2, ws.max_row + 1):
                current_id = str(ws.cell(row=row_number, column=col_map.get("proxy_id", 3)).value or "").strip()
                current_raw = str(ws.cell(row=row_number, column=col_map["proxy"]).value or "").strip()
                if (proxy_id and current_id == proxy_id) or (sheet_row == row_number) or (not proxy_id and raw and current_raw == raw):
                    ws.cell(row=row_number, column=col_map["status"], value=str(status or "INACTIVE").upper())
                    matched = True
                    break
            if matched:
                self._atomic_save(wb)
            wb.close()
            return matched

    def load_permanent_counts(self, proxy_pool) -> None:
        if not hasattr(proxy_pool, "set_permanent_count"):
            return
        with self._lock:
            wb = self._load_workbook(read_only=True)
            if "Success" not in wb.sheetnames:
                wb.close()
                return
            ws = wb["Success"]
            headers = self._headers(ws)
            col_map = self._col_map(headers)
            proxy_col = col_map.get("proxy_used")
            proxy_id_col = col_map.get("proxy_id")
            if not proxy_col:
                wb.close()
                return
            counts = {}
            for row in ws.iter_rows(min_row=2, values_only=True):
                proxy = str(row[proxy_col - 1] or "").strip()
                proxy_id = str(row[proxy_id_col - 1] or "").strip() if proxy_id_col else ""
                identity = proxy_id or proxy
                if identity and proxy.lower() != "direct":
                    counts[(identity, proxy)] = counts.get((identity, proxy), 0) + 1
            wb.close()
        for (identity, proxy), count in counts.items():
            proxy_pool.set_permanent_count(
                {"proxy_id": identity if identity.startswith("PX-") else "", "raw": proxy},
                count,
            )

    @staticmethod
    def success_stamp() -> str:
        return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def create_template(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Accounts"
    ws.append(ACCOUNT_HEADERS)
    ws.append(
        [
            "09000000000",
            "Example1234@",
            "ヤマダ",
            "タロウ",
            "2000-01-01",
            "1234",
            "PENDING",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
            "",
        ]
    )
    XlsxConnection._format_header(ws)
    ws_success = wb.create_sheet("Success")
    ws_success.append(SUCCESS_HEADERS)
    XlsxConnection._format_header(ws_success)
    ws_proxies = wb.create_sheet("Proxies")
    ws_proxies.append(PROXIES_HEADERS)
    ws_proxies.append(["host:port:user:pass", "DISABLED", "PX-000001"])
    XlsxConnection._format_header(ws_proxies)
    widths = {
        "A": 14,
        "B": 18,
        "C": 20,
        "D": 20,
        "E": 14,
        "F": 10,
        "G": 14,
        "H": 22,
        "I": 42,
        "J": 28,
        "K": 16,
        "L": 24,
        "M": 18,
        "N": 14,
        "O": 20,
    }
    for sheet in (ws, ws_success):
        for col, width in widths.items():
            sheet.column_dimensions[col].width = width
    ws_proxies.column_dimensions["A"].width = 40
    ws_proxies.column_dimensions["B"].width = 12
    ws_proxies.column_dimensions["C"].width = 16
    wb.save(path)
